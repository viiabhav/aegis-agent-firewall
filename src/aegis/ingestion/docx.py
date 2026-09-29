from __future__ import annotations

from pathlib import Path
from docx import Document


def extract_docx_text(path: Path) -> tuple[str, dict]:
    document = Document(str(path))
    parts: list[str] = []

    parts.extend(p.text for p in document.paragraphs if p.text.strip())

    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(f"[table:{table_index}:row:{row_index}] " + " | ".join(cells))

    for section_index, section in enumerate(document.sections, start=1):
        header = "\n".join(p.text for p in section.header.paragraphs if p.text.strip())
        footer = "\n".join(p.text for p in section.footer.paragraphs if p.text.strip())
        if header:
            parts.append(f"[header:{section_index}] {header}")
        if footer:
            parts.append(f"[footer:{section_index}] {footer}")

    return "\n".join(parts), {"filename": path.name, "parser": "python-docx"}
